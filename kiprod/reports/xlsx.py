from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from kiprod.analysis.comparison import Comparison
from kiprod.analysis.landed import LandedCost


def comparison_workbook(comparison: Comparison, costs: list[LandedCost], output: Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Comparison"
    summary.append(
        ["Rank", "Supplier", "Boards", "Currency", "Landed total", "Per board", "Quoted lines"]
    )
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for row in comparison.rows:
        summary.append(
            [
                row.rank,
                row.supplier_id,
                row.board_quantity,
                row.currency,
                float(row.total),
                float(row.per_board),
                row.quoted_lines,
            ]
        )
    details = workbook.create_sheet("Lines")
    details.append(
        [
            "Supplier",
            "Boards",
            "MPN",
            "Supplier P/N",
            "Required qty",
            "Buy qty",
            "Unit price",
            "Extended price",
            "Currency",
            "Source",
            "Source ref",
        ]
    )
    for cell in details[1]:
        cell.font = Font(bold=True)
    for cost in costs:
        for line in cost.lines:
            details.append(
                [
                    cost.supplier_id,
                    cost.board_quantity,
                    line.manufacturer_part_number,
                    line.supplier_part_number,
                    line.required_quantity,
                    line.buy_quantity,
                    float(line.unit_price),
                    float(line.extended_price),
                    line.currency,
                    line.source_ref.path,
                    line.source_ref.reference or "",
                ]
            )
    gaps = workbook.create_sheet("Coverage gaps")
    gaps.append(["MPN not offered"])
    gaps["A1"].font = Font(bold=True)
    for mpn in comparison.coverage_gaps:
        gaps.append([mpn])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
