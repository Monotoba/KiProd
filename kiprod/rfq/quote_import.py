from __future__ import annotations

import csv
import math
import re
import shutil
import subprocess
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import cast

from pydantic import BaseModel, Field

from kiprod.core.models import Offer, PriceBreak, SourceRef

_ALIASES: dict[str, set[str]] = {
    "mpn": {"mpn", "manufacturerpartnumber", "manufacturerpn"},
    "spn": {"supplierpartnumber", "supplierpn", "suppliernumber", "sku", "partnumber"},
    "quantity": {"quantity", "qty", "quantitybreak", "qtybreak"},
    "unit_price": {"unitprice", "priceeach", "priceperunit", "unitcost"},
    "currency": {"currency", "ccy"},
    "stock": {"stock", "stockquantity", "availablequantity"},
    "lead_time_days": {"leadtimedays", "leaddays", "leadtime"},
}
_REQUIRED = {"mpn", "spn", "quantity", "unit_price", "currency"}
_CONFIDENCE_THRESHOLD = 0.85


class QuoteReviewItem(BaseModel):
    source_ref: SourceRef
    confidence: float
    issue: str
    candidate_mpn: str | None = None


class ExtractedOffer(BaseModel):
    offer: Offer
    confidence: float
    requires_review: bool


class QuoteImportResult(BaseModel):
    supplier_id: str
    source_file: str
    offers: list[ExtractedOffer] = Field(default_factory=list)
    review_queue: list[QuoteReviewItem] = Field(default_factory=list)


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _header_map(headers: list[object]) -> tuple[dict[str, int], list[str]]:
    mapping: dict[str, int] = {}
    issues: list[str] = []
    reverse = {alias: key for key, aliases in _ALIASES.items() for alias in aliases}
    for index, header in enumerate(headers):
        key = reverse.get(_header_key(header))
        if key is None:
            continue
        if key in mapping:
            issues.append(f"multiple columns map to {key}")
        else:
            mapping[key] = index
    missing = sorted(_REQUIRED - mapping.keys())
    if missing:
        issues.append("missing required columns: " + ", ".join(missing))
    return mapping, issues


def _read_csv(path: Path) -> tuple[list[object], list[list[object]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return [], []
    return cast(tuple[list[object], list[list[object]]], (rows[0], rows[1:]))


def _read_xlsx(path: Path) -> tuple[list[object], list[list[object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX quote import requires `pip install kiprod[reports]`") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return [], []
    return list(rows[0]), [list(row) for row in rows[1:]]


def _read_pdf(path: Path) -> tuple[list[object], list[list[object]]]:
    executable = shutil.which("pdftotext")
    if executable is None:
        raise ValueError("PDF quote import requires the local `pdftotext` executable")
    try:
        result = subprocess.run(
            [executable, "-layout", "-enc", "UTF-8", str(path), "-"],
            capture_output=True,
            check=True,
            text=True,
            timeout=30,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        raise ValueError(f"could not extract quote text from PDF {path}: {exc}") from exc
    reverse = {alias: key for key, aliases in _ALIASES.items() for alias in aliases}
    headers = [
        "MPN",
        "Supplier Part Number",
        "Quantity",
        "Unit Price",
        "Currency",
        "Stock",
        "Lead Time Days",
    ]
    header_keys = [reverse[_header_key(header)] for header in headers]
    blocks = re.split(r"\n\s*\n|\f", result.stdout)
    rows: list[list[object]] = []
    for block in blocks:
        fields: dict[str, str] = {}
        for line in block.splitlines():
            label, separator, value = line.partition(":")
            if not separator:
                continue
            key = reverse.get(_header_key(label))
            if key is None:
                continue
            cleaned = value.strip()
            if key in fields and fields[key] != cleaned:
                fields[key] = ""
            else:
                fields[key] = cleaned
        if fields:
            rows.append([fields.get(key, "") for key in header_keys])
    return cast(tuple[list[object], list[list[object]]], (headers, rows))


def _cell(row: list[object], mapping: dict[str, int], name: str) -> str:
    index = mapping.get(name)
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _positive_int(value: str, field: str) -> int:
    number = Decimal(value)
    if number != number.to_integral_value() or number <= 0:
        raise ValueError(f"{field} must be a positive integer")
    return int(number)


def _optional_int(value: str, field: str) -> int | None:
    if not value:
        return None
    number = Decimal(value)
    if number != number.to_integral_value() or number < 0:
        raise ValueError(f"{field} must be a non-negative integer")
    return int(number)


def _row_value(row: list[object], mapping: dict[str, int]) -> dict[str, str]:
    return {name: _cell(row, mapping, name) for name in mapping}


def parse_quote_attachment(
    path: Path,
    supplier_id: str,
    *,
    currency: str | None = None,
) -> QuoteImportResult:
    """Parse CSV/XLSX tables or labeled PDF blocks into offers and review items."""
    suffix = path.suffix.casefold()
    source_kind = "row"
    if suffix == ".csv":
        headers, rows = _read_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        headers, rows = _read_xlsx(path)
    elif suffix == ".pdf":
        headers, rows = _read_pdf(path)
        source_kind = "block"
    else:
        raise ValueError(f"unsupported quote attachment format: {suffix or 'unknown'}")
    if not supplier_id.strip():
        raise ValueError("supplier ID must be non-empty")
    mapping, header_issues = _header_map(headers)
    if currency is not None:
        normalized_currency = currency.strip().upper()
        if not re.fullmatch(r"[A-Z]{3}", normalized_currency):
            raise ValueError("currency must be a three-letter code")
        if "currency" not in mapping:
            mapping["currency"] = -1
            header_issues = [
                issue for issue in header_issues if issue.startswith("multiple columns")
            ]
            missing = sorted(_REQUIRED - mapping.keys())
            if missing:
                header_issues.append("missing required columns: " + ", ".join(missing))
    supplier_id = supplier_id.strip()
    source_path = path.as_posix() if not path.is_absolute() else path.name
    result = QuoteImportResult(supplier_id=supplier_id, source_file=source_path)
    if header_issues:
        result.review_queue.append(
            QuoteReviewItem(
                source_ref=SourceRef(path=source_path, reference="header"),
                confidence=0.2,
                issue="; ".join(header_issues),
            )
        )
        return result
    if suffix == ".pdf" and not rows:
        result.review_queue.append(
            QuoteReviewItem(
                source_ref=SourceRef(path=source_path, reference="pdf"),
                confidence=0.1,
                issue="no labeled quote fields could be extracted from PDF text",
            )
        )
        return result

    grouped: dict[tuple[str, str, str], list[tuple[int, PriceBreak, int | None, int | None]]] = (
        defaultdict(list)
    )
    bad_keys: set[tuple[str, str, str]] = set()
    for row_number, row in enumerate(rows, start=1 if source_kind == "block" else 2):
        if not any(str(value or "").strip() for value in row):
            continue
        values = _row_value(row, mapping)
        if mapping.get("currency") == -1 or (not values.get("currency") and currency is not None):
            values["currency"] = normalized_currency
        reference = SourceRef(path=source_path, reference=f"{source_kind}:{row_number}")
        mpn = values.get("mpn", "")
        spn = values.get("spn", "")
        key: tuple[str, str, str] | None = None
        try:
            if not mpn or not spn:
                raise ValueError("manufacturer and supplier part numbers are required")
            row_currency = values.get("currency", "").upper()
            if not re.fullmatch(r"[A-Z]{3}", row_currency):
                raise ValueError("currency must be a three-letter code")
            key = (mpn, spn, row_currency)
            quantity = _positive_int(values.get("quantity", ""), "quantity")
            unit_price = Decimal(values.get("unit_price", ""))
            if not unit_price.is_finite() or unit_price < 0:
                raise ValueError("unit price must be a finite non-negative amount")
            numeric_price = float(unit_price)
            if not math.isfinite(numeric_price):
                raise ValueError("unit price is outside the supported numeric range")
            stock = _optional_int(values.get("stock", ""), "stock")
            lead_time = _optional_int(values.get("lead_time_days", ""), "lead time")
            price = PriceBreak(
                quantity=quantity,
                unit_price=numeric_price,
                currency=row_currency,
            )
            if any(item[1].quantity == quantity for item in grouped[key]):
                bad_keys.add(key)
                previous = next(item[1] for item in grouped[key] if item[1].quantity == quantity)
                if previous.unit_price != price.unit_price:
                    bad_keys.add(key)
                    raise ValueError("conflicting prices for the same quantity break")
                raise ValueError("duplicate quantity break")
            grouped[key].append((row_number, price, stock, lead_time))
        except (InvalidOperation, ValueError) as exc:
            if key is not None:
                bad_keys.add(key)
            result.review_queue.append(
                QuoteReviewItem(
                    source_ref=reference,
                    confidence=0.35,
                    issue=str(exc),
                    candidate_mpn=mpn or None,
                )
            )

    for key, items in sorted(grouped.items()):
        if key in bad_keys:
            continue
        mpn, spn, _ = key
        rows_ref = ",".join(str(item[0]) for item in items)
        offer = Offer(
            supplier_id=supplier_id,
            supplier_part_number=spn,
            manufacturer_part_number=mpn,
            price_breaks=sorted((item[1] for item in items), key=lambda item: item.quantity),
            stock=next((item[2] for item in items if item[2] is not None), None),
            lead_time_days=next((item[3] for item in items if item[3] is not None), None),
            source_ref=SourceRef(path=source_path, reference=f"{source_kind}s:{rows_ref}"),
        )
        confidence = 0.72 if suffix == ".pdf" else (0.97 if suffix == ".csv" else 0.95)
        result.offers.append(
            ExtractedOffer(
                offer=offer,
                confidence=confidence,
                requires_review=confidence < _CONFIDENCE_THRESHOLD,
            )
        )
        if confidence < _CONFIDENCE_THRESHOLD:
            result.review_queue.append(
                QuoteReviewItem(
                    source_ref=offer.source_ref,
                    confidence=confidence,
                    issue="PDF-extracted quote requires human verification",
                    candidate_mpn=mpn,
                )
            )
    return result
